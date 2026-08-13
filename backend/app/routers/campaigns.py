import csv
import io
import logging
import re

from fastapi import APIRouter, Depends, HTTPException, UploadFile, File, Query
from fastapi.responses import StreamingResponse, Response
from sqlalchemy.orm import Session

from ..database import get_db
from ..models import Campaign, Project, Recipient, CampaignContent, CampaignAttachment, Message, User, Template
from ..schemas import (
    CampaignOut, CampaignCreate, CampaignUpdate, ContentIn, ContentOut,
    RecipientOut, RecipientIn, RecipientActiveIn, PreviewIn, PreviewOut,
    TemplatePreviewIn, TemplatePreviewOut, CampaignAttachmentOut,
    DatasetValidationReport,
    CampaignReviewOut, CampaignReviewPreviewIn,
)
from ..deps import require, get_current_user, user_permissions
from ..campaign_utils import (
    render_template as _render, valid_email as _valid_email, valid_mobile as _valid_mobile,
)
from ..data_validator import parse_upload, validate_rows, sample_csv_bytes, error_report_csv
from ..email_template import render_email_template
from ..storage import save_campaign_attachment, delete_campaign_attachment
from ..worker import enqueue_message
from ..campaign_reviewer import review_campaign
from ..ai_content import ai_review_campaign, ai_score_dataset

router = APIRouter(prefix="/api", tags=["campaigns"])
logger = logging.getLogger(__name__)

# Content is meant to be plain text with {{Placeholder}} tokens -- a since-reverted
# HTML/signature feature let raw markup leak into saved content, so reject it here.
HTML_TAG_RE = re.compile(r"<[a-zA-Z][^>]*>")


# ---------- helpers ----------
def _accessible_project(db: Session, project_id: int, user: User) -> Project:
    project = db.get(Project, project_id)
    if not project:
        raise HTTPException(status_code=404, detail="Project not found")
    perms = user_permissions(user)
    see_all = "user.manage" in perms or "system.configure" in perms
    if not see_all and project.id not in {p.id for p in user.projects}:
        raise HTTPException(status_code=403, detail="No access to this project")
    return project


def _get_campaign(db: Session, campaign_id: int, user: User) -> Campaign:
    campaign = db.get(Campaign, campaign_id)
    if not campaign:
        raise HTTPException(status_code=404, detail="Campaign not found")
    _accessible_project(db, campaign.project_id, user)
    return campaign


def _columns(campaign: Campaign) -> list[str]:
    cols: list[str] = []
    for r in campaign.recipients:
        for k in (r.data or {}).keys():
            if k not in cols:
                cols.append(k)
    return cols


def _to_out(campaign: Campaign) -> CampaignOut:
    out = CampaignOut.model_validate(campaign)
    out.recipient_count = len(campaign.recipients)
    out.columns = _columns(campaign)
    out.project_name = campaign.project.name if campaign.project else ""
    return out


# ---------- campaign CRUD ----------
@router.get("/campaigns", response_model=list[CampaignOut])
def list_all_campaigns(db: Session = Depends(get_db), user: User = Depends(get_current_user)):
    """All campaigns across the projects the user can see."""
    perms = user_permissions(user)
    see_all = "user.manage" in perms or "system.configure" in perms
    if see_all:
        pids = [pid for (pid,) in db.query(Project.id).all()]
    else:
        pids = [p.id for p in user.projects]
    if not pids:
        return []
    campaigns = db.query(Campaign).filter(Campaign.project_id.in_(pids)).order_by(Campaign.id.desc()).all()
    return [_to_out(c) for c in campaigns]


@router.get("/projects/{project_id}/campaigns", response_model=list[CampaignOut])
def list_campaigns(project_id: int, db: Session = Depends(get_db), user: User = Depends(get_current_user)):
    _accessible_project(db, project_id, user)
    campaigns = db.query(Campaign).filter_by(project_id=project_id).order_by(Campaign.id.desc()).all()
    return [_to_out(c) for c in campaigns]


@router.get("/campaigns/{campaign_id}", response_model=CampaignOut)
def get_campaign(campaign_id: int, db: Session = Depends(get_db), user: User = Depends(get_current_user)):
    return _to_out(_get_campaign(db, campaign_id, user))


@router.post("/campaigns", response_model=CampaignOut)
def create_campaign(payload: CampaignCreate, db: Session = Depends(get_db),
                    user: User = Depends(require("campaign.create"))):
    _accessible_project(db, payload.project_id, user)
    campaign = Campaign(project_id=payload.project_id, name=payload.name,
                        description=payload.description, created_by=user.id)
    db.add(campaign)
    db.commit()
    db.refresh(campaign)
    return _to_out(campaign)


@router.put("/campaigns/{campaign_id}", response_model=CampaignOut)
def update_campaign(campaign_id: int, payload: CampaignUpdate, db: Session = Depends(get_db),
                    user: User = Depends(require("campaign.edit"))):
    campaign = _get_campaign(db, campaign_id, user)
    for k, v in payload.model_dump(exclude_unset=True).items():
        setattr(campaign, k, v)
    db.commit()
    db.refresh(campaign)
    return _to_out(campaign)




# ---------- AI Smart Data Validation ----------
@router.get("/campaigns/{campaign_id}/dataset/sample.csv")
def download_sample_csv(campaign_id: int, db: Session = Depends(get_db),
                        user: User = Depends(get_current_user)):
    """Sample CSV with required + optional columns for audience import."""
    _get_campaign(db, campaign_id, user)
    data = sample_csv_bytes()
    return Response(
        content=data,
        media_type="text/csv",
        headers={"Content-Disposition": 'attachment; filename="reach_contacts_sample.csv"'},
    )


@router.post("/campaigns/{campaign_id}/dataset/validate", response_model=DatasetValidationReport)
async def validate_dataset(
    campaign_id: int,
    file: UploadFile = File(...),
    replace: bool = Query(True, description="If true, ignore existing campaign contacts when flagging duplicates"),
    db: Session = Depends(get_db),
    user: User = Depends(require("dataset.upload")),
):
    """Run AI Smart Validation without importing. Returns quality score + row issues."""
    campaign = _get_campaign(db, campaign_id, user)
    raw = await file.read()
    try:
        headers, rows = parse_upload(raw, file.filename or "", file.content_type or "")
    except ValueError as e:
        raise HTTPException(status_code=400, detail=str(e))

    # When replacing the list, do not mark rows invalid only because the email
    # already exists on this campaign.
    existing_emails: set[str] = set()
    existing_mobiles: set[str] = set()
    if not replace:
        existing_emails = {(r.email or "").lower() for r in campaign.recipients if r.email}
        for r in campaign.recipients:
            if r.mobile:
                existing_mobiles.add(re.sub(r"\D", "", r.mobile))

    report = validate_rows(
        headers, rows,
        existing_emails=existing_emails,
        existing_mobiles=existing_mobiles,
    )

    # AI quality score when OpenAI is configured; keep rule metrics as facts
    ai = ai_score_dataset(report)
    if ai:
        report["quality_score"] = ai["quality_score"]
        # Prefer AI recommendations; keep rule tips that are not duplicated
        merged = list(ai.get("recommendations") or [])
        for tip in (report.get("recommendations") or []):
            if tip and tip not in merged:
                merged.append(tip)
        report["recommendations"] = merged[:8]
        if ai.get("summary_note"):
            report["recommendations"] = [ai["summary_note"]] + [
                r for r in report["recommendations"] if r != ai["summary_note"]
            ][:7]
        report["score_source"] = "openai"
        print(f"[dataset] AI quality_score={ai['quality_score']}")
    else:
        report["score_source"] = "rules"
        print("[dataset] AI score unavailable — using rule quality_score")

    out_rows = report["rows"]
    if len(out_rows) > 500:
        problems = [r for r in out_rows if r["status"] != "valid"]
        valids = [r for r in out_rows if r["status"] == "valid"][:50]
        out_rows = problems + valids

    return DatasetValidationReport(
        quality_score=report["quality_score"],
        score_breakdown=report["score_breakdown"],
        summary=report["summary"],
        column_mapping=report["column_mapping"],
        recommendations=report["recommendations"],
        rows=out_rows,
        importable_count=report["summary"]["ready_for_import"],
    )


@router.post("/campaigns/{campaign_id}/dataset/validate/report.csv")
async def download_validation_report(campaign_id: int, file: UploadFile = File(...),
                                     db: Session = Depends(get_db),
                                     user: User = Depends(require("dataset.upload"))):
    """Re-validate and return a CSV error/warning report for download."""
    campaign = _get_campaign(db, campaign_id, user)
    raw = await file.read()
    try:
        headers, rows = parse_upload(raw, file.filename or "", file.content_type or "")
    except ValueError as e:
        raise HTTPException(status_code=400, detail=str(e))
    existing_emails = { (r.email or "").lower() for r in campaign.recipients if r.email }
    report = validate_rows(headers, rows, existing_emails=existing_emails)
    data = error_report_csv(report)
    return Response(
        content=data,
        media_type="text/csv",
        headers={"Content-Disposition": 'attachment; filename="validation_report.csv"'},
    )


@router.post("/campaigns/{campaign_id}/dataset/import-valid", response_model=CampaignOut)
async def import_valid_dataset(
    campaign_id: int,
    file: UploadFile = File(...),
    include_warnings: bool = Query(True, description="Import rows with warnings as well as fully valid rows"),
    replace: bool = Query(True, description="Replace existing recipients"),
    db: Session = Depends(get_db),
    user: User = Depends(require("dataset.upload")),
):
    """Validate then import only ready rows. Default: replace existing recipients."""
    campaign = _get_campaign(db, campaign_id, user)
    raw = await file.read()
    if not raw:
        raise HTTPException(status_code=400, detail="Empty file — nothing to import")
    try:
        headers, rows = parse_upload(raw, file.filename or "", file.content_type or "")
    except ValueError as e:
        raise HTTPException(status_code=400, detail=str(e))

    # Do not treat current campaign contacts as duplicates — this upload replaces them
    report = validate_rows(headers, rows, existing_emails=set(), existing_mobiles=set())
    allowed = {"valid"} if not include_warnings else {"valid", "warning"}
    importable = [r for r in report["rows"] if r["status"] in allowed]
    if not importable:
        raise HTTPException(
            status_code=400,
            detail=(
                "No valid rows available to import. Fix invalid emails/phones "
                f"(invalid={report['summary'].get('invalid_rows', 0)}) or re-upload a corrected file."
            ),
        )

    mapping = report["column_mapping"]
    name_col, email_col, mobile_col = mapping["name"], mapping["email"], mapping["mobile"]
    if not (name_col and email_col and mobile_col):
        raise HTTPException(status_code=400, detail="Missing mandatory column mapping (Name, Email, Mobile)")

    if replace:
        db.query(Recipient).filter_by(campaign_id=campaign.id).delete()

    for r in importable:
        data = dict(r["data"] or {})
        data["Name"] = r["name"]
        data["Email"] = r["email"]
        data["Mobile"] = r["mobile"]
        rec = Recipient(
            campaign_id=campaign.id,
            name=r["name"],
            email=r["email"],
            mobile=r["mobile"],
            data=data,
            active=True,
        )
        db.add(rec)
    db.commit()
    db.refresh(campaign)
    return _to_out(campaign)


# ---------- dataset ----------
@router.post("/campaigns/{campaign_id}/dataset", response_model=CampaignOut)
async def upload_dataset(campaign_id: int, file: UploadFile = File(...), db: Session = Depends(get_db),
                         user: User = Depends(require("dataset.upload"))):
    campaign = _get_campaign(db, campaign_id, user)
    raw = await file.read()
    filename = (file.filename or "").lower()

    rows: list[dict] = []
    if filename.endswith(".csv") or file.content_type in ("text/csv", "application/vnd.ms-excel"):
        text = raw.decode("utf-8-sig", errors="replace")
        reader = csv.DictReader(io.StringIO(text))
        rows = [dict(r) for r in reader]
    elif filename.endswith(".xlsx"):
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        ws = wb.active
        it = ws.iter_rows(values_only=True)
        headers = [str(h).strip() if h is not None else "" for h in next(it, [])]
        for values in it:
            if values is None:
                continue
            row = {}
            for i, h in enumerate(headers):
                if not h:
                    continue
                val = values[i] if i < len(values) else None
                row[h] = "" if val is None else str(val)
            if any(v for v in row.values()):
                rows.append(row)
    else:
        raise HTTPException(status_code=400, detail="Unsupported file type. Upload .csv or .xlsx")

    if not rows:
        raise HTTPException(status_code=400, detail="No data rows found in the file")

    # validate mandatory columns (case-insensitive)
    headers_lower = {k.lower(): k for k in rows[0].keys()}

    def find(*names):
        for n in names:
            if n in headers_lower:
                return headers_lower[n]
        return None

    name_col = find("name")
    email_col = find("email", "email address", "email id")
    mobile_col = find("mobile", "mobile number", "phone", "mobile no")
    missing = [label for label, col in
               [("Name", name_col), ("Email Address", email_col), ("Mobile Number", mobile_col)] if col is None]
    if missing:
        raise HTTPException(status_code=400, detail=f"Missing mandatory column(s): {', '.join(missing)}")

    # replace existing recipients
    db.query(Recipient).filter_by(campaign_id=campaign.id).delete()
    imported = 0
    for row in rows:
        clean = {k.strip(): (v.strip() if isinstance(v, str) else v) for k, v in row.items() if k}
        rec = Recipient(
            campaign_id=campaign.id,
            name=str(clean.get(name_col, "")).strip(),
            email=str(clean.get(email_col, "")).strip(),
            mobile=str(clean.get(mobile_col, "")).strip(),
            data=clean,
        )
        db.add(rec)
        imported += 1
    db.commit()
    db.refresh(campaign)
    return _to_out(campaign)


@router.get("/campaigns/{campaign_id}/recipients", response_model=list[RecipientOut])
def list_recipients(campaign_id: int, db: Session = Depends(get_db), user: User = Depends(get_current_user)):
    campaign = _get_campaign(db, campaign_id, user)
    return campaign.recipients


def _recipient_data(campaign: Campaign, name: str, email: str, mobile: str) -> dict:
    """Build placeholder data, reusing the campaign's existing column keys when present."""
    existing = list(campaign.recipients[0].data.keys()) if campaign.recipients else []

    def find(*cands):
        for k in existing:
            if k.lower() in cands:
                return k
        return None

    name_k = find("name") or "Name"
    email_k = find("email", "email address", "email id") or "Email Address"
    mobile_k = find("mobile", "mobile number", "phone", "mobile no") or "Mobile Number"
    data = {name_k: name, email_k: email, mobile_k: mobile}
    for k in existing:
        data.setdefault(k, "")
    return data


@router.post("/campaigns/{campaign_id}/recipients", response_model=RecipientOut)
def add_recipient(campaign_id: int, payload: RecipientIn, db: Session = Depends(get_db),
                  user: User = Depends(require("dataset.upload"))):
    campaign = _get_campaign(db, campaign_id, user)
    name, email, mobile = payload.name.strip(), payload.email.strip(), payload.mobile.strip()
    if not name:
        raise HTTPException(status_code=400, detail="Name is required")
    if not email and not mobile:
        raise HTTPException(status_code=400, detail="Provide an email address or a mobile number")
    rec = Recipient(campaign_id=campaign.id, name=name, email=email, mobile=mobile,
                    data=_recipient_data(campaign, name, email, mobile))
    db.add(rec)
    db.commit()
    db.refresh(rec)
    return rec


@router.patch("/campaigns/{campaign_id}/recipients/{recipient_id}", response_model=RecipientOut)
def set_recipient_active(campaign_id: int, recipient_id: int, payload: RecipientActiveIn,
                         db: Session = Depends(get_db), user: User = Depends(require("dataset.upload"))):
    campaign = _get_campaign(db, campaign_id, user)
    rec = db.get(Recipient, recipient_id)
    if not rec or rec.campaign_id != campaign.id:
        raise HTTPException(status_code=404, detail="Recipient not found")
    rec.active = payload.active
    db.commit()
    db.refresh(rec)
    return rec


@router.delete("/campaigns/{campaign_id}/recipients/{recipient_id}")
def remove_recipient(campaign_id: int, recipient_id: int, db: Session = Depends(get_db),
                     user: User = Depends(require("dataset.upload"))):
    campaign = _get_campaign(db, campaign_id, user)
    rec = db.get(Recipient, recipient_id)
    if not rec or rec.campaign_id != campaign.id:
        raise HTTPException(status_code=404, detail="Recipient not found")
    db.delete(rec)
    db.commit()
    return {"ok": True}


# ---------- content ----------
@router.get("/campaigns/{campaign_id}/content", response_model=list[ContentOut])
def get_content(campaign_id: int, db: Session = Depends(get_db), user: User = Depends(get_current_user)):
    campaign = _get_campaign(db, campaign_id, user)
    return campaign.contents


@router.put("/campaigns/{campaign_id}/content/{channel}", response_model=ContentOut)
def set_content(campaign_id: int, channel: str, payload: ContentIn, db: Session = Depends(get_db),
                user: User = Depends(require("content.edit"))):
    if channel not in ("email", "whatsapp", "sms"):
        raise HTTPException(status_code=400, detail="Invalid channel")
    # Only the email channel supports the branded-template mode; whatsapp/sms
    # always stay plain text regardless of what the client sends.
    content_mode = payload.content_mode if channel == "email" else "plain"
    if content_mode not in ("plain", "template"):
        raise HTTPException(status_code=400, detail="Invalid content mode")
    # The subject line is always plain text, in both modes. The body is only
    # checked when it's actually going to be used (plain mode) -- in template
    # mode the structured fields are the source of truth and body is ignored.
    if channel == "email" and HTML_TAG_RE.search(payload.subject):
        raise HTTPException(
            status_code=400,
            detail="Email subject must be plain text -- HTML markup is not allowed",
        )
    if channel == "email" and content_mode == "plain" and HTML_TAG_RE.search(payload.body):
        raise HTTPException(
            status_code=400,
            detail="Email subject/body must be plain text -- HTML markup is not allowed",
        )
    campaign = _get_campaign(db, campaign_id, user)
    content = db.query(CampaignContent).filter_by(campaign_id=campaign.id, channel=channel).first()
    if not content:
        content = CampaignContent(campaign_id=campaign.id, channel=channel)
        db.add(content)
    content.subject = payload.subject
    content.content_mode = content_mode
    if content_mode == "template":
        content.body = ""
        content.template_fields = payload.template_fields.model_dump()
    else:
        content.body = payload.body
        content.template_fields = {}
    db.commit()
    db.refresh(content)
    return content


@router.post("/campaigns/{campaign_id}/content/{channel}/from-template/{template_id}", response_model=ContentOut)
def set_content_from_template(campaign_id: int, channel: str, template_id: int, db: Session = Depends(get_db),
                              user: User = Depends(require("content.edit"))):
    """Copy a library template's resolved content into this campaign's own
    content storage. From this point on it's an independent copy -- editing
    it here never touches the source `templates` row, and editing the source
    template later never changes what this campaign is shown to have sent
    (source_template_id is kept only so reporting can trace it back)."""
    if channel not in ("email", "whatsapp"):
        raise HTTPException(status_code=400, detail="Invalid channel")
    campaign = _get_campaign(db, campaign_id, user)
    template = db.get(Template, template_id)
    if not template or template.channel != channel:
        raise HTTPException(status_code=404, detail="Template not found")
    if template.status != "published":
        raise HTTPException(status_code=400, detail="Only published templates can be used in a campaign")

    content = db.query(CampaignContent).filter_by(campaign_id=campaign.id, channel=channel).first()
    if not content:
        content = CampaignContent(campaign_id=campaign.id, channel=channel)
        db.add(content)

    if channel == "email":
        if not template.email_content:
            raise HTTPException(status_code=400, detail="This template has no email content")
        ec = template.email_content
        fields = dict(ec.fields or {})
        branding = dict(ec.branding_override or {}) if isinstance(ec.branding_override, dict) else {}
        # Templates saved from campaign/AI plain copy → load as plain text in Content editor
        if branding.get("source") == "campaign_plain":
            content.subject = branding.get("plain_subject") or ec.subject or ""
            content.body = branding.get("plain_body") or fields.get("opening_line") or ""
            content.content_mode = "plain"
            content.template_fields = {}
        else:
            content.subject = ec.subject or ""
            content.content_mode = "template"
            content.body = ""
            content.template_fields = fields
    else:  # whatsapp
        if not template.whatsapp_content:
            raise HTTPException(status_code=400, detail="This template has no WhatsApp content")
        wc = template.whatsapp_content
        content.subject = ""
        content.content_mode = "plain"
        content.body = wc.body_text
        content.template_fields = {}

    content.source_template_id = template.id
    db.commit()
    db.refresh(content)
    return content


# ---------- attachments ----------
@router.get("/campaigns/{campaign_id}/attachments", response_model=list[CampaignAttachmentOut])
def list_attachments(campaign_id: int, db: Session = Depends(get_db), user: User = Depends(get_current_user)):
    campaign = _get_campaign(db, campaign_id, user)
    return campaign.attachments


@router.post("/campaigns/{campaign_id}/attachments", response_model=CampaignAttachmentOut)
async def upload_attachment(campaign_id: int, file: UploadFile = File(...), db: Session = Depends(get_db),
                            user: User = Depends(require("content.edit"))):
    campaign = _get_campaign(db, campaign_id, user)
    existing_total = sum(a.size_bytes for a in campaign.attachments)
    storage_path, size_bytes, content_type = await save_campaign_attachment(file, existing_total)
    att = CampaignAttachment(
        campaign_id=campaign.id, filename=file.filename or "attachment",
        content_type=content_type, storage_path=storage_path, size_bytes=size_bytes,
    )
    db.add(att)
    db.commit()
    db.refresh(att)
    return att


@router.delete("/campaigns/{campaign_id}/attachments/{attachment_id}")
def remove_attachment(campaign_id: int, attachment_id: int, db: Session = Depends(get_db),
                      user: User = Depends(require("content.edit"))):
    campaign = _get_campaign(db, campaign_id, user)
    att = db.get(CampaignAttachment, attachment_id)
    if not att or att.campaign_id != campaign.id:
        raise HTTPException(status_code=404, detail="Attachment not found")
    delete_campaign_attachment(att.storage_path)
    db.delete(att)
    db.commit()
    return {"ok": True}


@router.post("/campaigns/{campaign_id}/preview", response_model=PreviewOut)
def preview(campaign_id: int, payload: PreviewIn, db: Session = Depends(get_db), user: User = Depends(get_current_user)):
    campaign = _get_campaign(db, campaign_id, user)
    recipient = None
    if payload.recipient_id:
        recipient = db.get(Recipient, payload.recipient_id)
        if recipient is None or recipient.campaign_id != campaign.id:
            raise HTTPException(status_code=404, detail="Recipient not found")
    if recipient is None and campaign.recipients:
        recipient = campaign.recipients[0]
    data = dict((recipient.data if recipient else {}) or {})
    missing: list[str] = []
    return PreviewOut(
        subject=_render(payload.subject, data, missing),
        body=_render(payload.body, data, missing),
        missing=sorted(set(missing)),
    )


@router.post("/campaigns/{campaign_id}/preview-template", response_model=TemplatePreviewOut)
def preview_template(campaign_id: int, payload: TemplatePreviewIn, db: Session = Depends(get_db),
                     user: User = Depends(get_current_user)):
    """Live preview for the branded-template email content mode -- renders the
    fixed HTML shell (never user-typed markup) with the campaign's unsaved
    structured field values, for the sandboxed iframe preview on the frontend."""
    campaign = _get_campaign(db, campaign_id, user)
    recipient = None
    if payload.recipient_id:
        recipient = db.get(Recipient, payload.recipient_id)
        if recipient is None or recipient.campaign_id != campaign.id:
            raise HTTPException(status_code=404, detail="Recipient not found")
    if recipient is None and campaign.recipients:
        recipient = campaign.recipients[0]
    missing: list[str] = []
    html_out = render_email_template(campaign.project, payload.fields.model_dump(), recipient, missing)
    return TemplatePreviewOut(html=html_out, missing=sorted(set(missing)))


# ---------- send ----------
@router.post("/campaigns/{campaign_id}/send", response_model=CampaignOut)
def send_campaign(campaign_id: int, db: Session = Depends(get_db), user: User = Depends(require("campaign.send"))):
    campaign = _get_campaign(db, campaign_id, user)
    project = campaign.project

    if campaign.status == "sending":
        raise HTTPException(status_code=409, detail="Campaign is currently sending, please wait")

    active_recipients = [r for r in campaign.recipients if r.active]
    if not active_recipients:
        raise HTTPException(status_code=400, detail="Upload a recipient dataset before sending")

    channels = []
    if campaign.email_enabled:
        channels.append("email")
    if campaign.whatsapp_enabled and project.whatsapp_active:
        channels.append("whatsapp")
    if campaign.sms_enabled and project.sms_active:
        channels.append("sms")
    if not channels:
        raise HTTPException(status_code=400, detail="No active channels for this campaign (check project channel config)")

    content_map = {c.channel: c for c in campaign.contents}
    for ch in channels:
        if ch not in content_map:
            raise HTTPException(status_code=400, detail=f"No content authored for channel: {ch}")

    # clear previous run
    db.query(Message).filter_by(campaign_id=campaign.id).delete()
    db.commit()

    queued_ids = []
    for rec in active_recipients:
        for ch in channels:
            to = rec.email if ch == "email" else rec.mobile
            valid = _valid_email(to) if ch == "email" else _valid_mobile(to)
            msg = Message(
                campaign_id=campaign.id, recipient_id=rec.id, channel=ch,
                to_address=to, recipient_name=rec.name,
            )
            if not valid:
                msg.status = "failed"
                msg.error = "Invalid email address" if ch == "email" else "Invalid mobile number"
            else:
                msg.status = "queued"
            db.add(msg)
            # Commit each row as soon as it's created so a mid-loop crash never
            # leaves a message unaccounted for.
            db.commit()
            if msg.status == "queued":
                queued_ids.append(msg.id)

    campaign.status = "sending"
    db.commit()
    db.refresh(campaign)

    # Enqueue only after every Message row is committed -- a worker picking up
    # the first job must always find the rest of the campaign's rows in place.
    for mid in queued_ids:
        enqueue_message(mid)

    # Auto-arm multi-channel sequence: schedule step 1 at its day + time
    if getattr(campaign, "followup_enabled", True):
        from datetime import datetime, timezone
        from ..models import CampaignSequenceStep
        from ..routers.engagement import compute_next_send_at
        step1 = (
            db.query(CampaignSequenceStep)
            .filter_by(campaign_id=campaign.id, step_number=1)
            .first()
        )
        if step1 is not None and step1.delay_days is not None:
            days = max(0, int(step1.delay_days))
        elif getattr(campaign, "followup_interval_days", None) is not None:
            days = max(0, int(campaign.followup_interval_days))
        else:
            days = 3
        send_time = (step1.send_time if step1 else None) or "10:00"
        now = datetime.now(timezone.utc)
        for rec in campaign.recipients:
            if rec.active and not getattr(rec, "campaign_stopped", False):
                if getattr(rec, "reply_status", "not_replied") == "not_replied":
                    rec.next_followup_at = compute_next_send_at(now, days, send_time)
                    rec.followup_count = 0
        db.commit()

    logger.info("campaign %s: send requested, %s message(s) enqueued across channels %s",
                campaign.id, len(queued_ids), channels)
    return _to_out(campaign)


# ---------- AI campaign review ----------
def _content_map(campaign: Campaign) -> dict[str, CampaignContent]:
    return {c.channel: c for c in (campaign.contents or [])}



def _smart_review(**kwargs) -> dict:
    """Prefer AI scoring; fall back to deterministic rules."""
    goal = kwargs.pop("goal", None) or kwargs.get("description") or ""
    ai = ai_review_campaign(
        name=kwargs.get("name") or "",
        description=kwargs.get("description") or "",
        goal=goal,
        email_enabled=bool(kwargs.get("email_enabled")),
        whatsapp_enabled=bool(kwargs.get("whatsapp_enabled")),
        sms_enabled=bool(kwargs.get("sms_enabled")),
        email_subject=kwargs.get("email_subject") or "",
        email_body=kwargs.get("email_body") or "",
        whatsapp_body=kwargs.get("whatsapp_body") or "",
        sms_body=kwargs.get("sms_body") or "",
    )
    if ai:
        return ai
    out = review_campaign(**kwargs)
    out["review_source"] = "rules"
    return out

@router.post("/campaigns/{campaign_id}/review", response_model=CampaignReviewOut)
def review_saved_campaign(campaign_id: int, db: Session = Depends(get_db),
                          user: User = Depends(get_current_user)):
    """Review a saved campaign's content before send. Does not modify the campaign."""
    campaign = _get_campaign(db, campaign_id, user)
    by_ch = _content_map(campaign)
    email = by_ch.get("email")
    wa = by_ch.get("whatsapp")
    sms = by_ch.get("sms")
    result = _smart_review(
        name=campaign.name or "",
        description=campaign.description or "",
        email_enabled=bool(campaign.email_enabled),
        whatsapp_enabled=bool(campaign.whatsapp_enabled),
        sms_enabled=bool(campaign.sms_enabled),
        email_subject=(email.subject if email else "") or "",
        email_body=(email.body if email else "") or "",
        email_content_mode=(email.content_mode if email else "plain") or "plain",
        email_template_fields=(email.template_fields if email else None) or {},
        whatsapp_body=(wa.body if wa else "") or "",
        sms_body=(sms.body if sms else "") or "",
    )
    return result


@router.post("/campaigns/review-preview", response_model=CampaignReviewOut)
def review_campaign_preview(payload: CampaignReviewPreviewIn,
                            user: User = Depends(get_current_user)):
    """Review draft content that may not be persisted yet (composer preview)."""
    fields = payload.email_template_fields.model_dump() if payload.email_template_fields else {}
    return _smart_review(
        name=payload.name,
        description=payload.description,
        email_enabled=payload.email_enabled,
        whatsapp_enabled=payload.whatsapp_enabled,
        sms_enabled=payload.sms_enabled,
        email_subject=payload.email_subject,
        email_body=payload.email_body,
        email_content_mode=payload.email_content_mode,
        email_template_fields=fields,
        whatsapp_body=payload.whatsapp_body,
        sms_body=payload.sms_body,
    )
